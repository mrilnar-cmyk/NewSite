from django.shortcuts import render, get_object_or_404, redirect
from django.http import JsonResponse
from django.views import View
from django.views.generic import ListView, DetailView, CreateView, UpdateView, DeleteView
from django.contrib.auth.mixins import LoginRequiredMixin
from django.urls import reverse_lazy
from django.contrib import messages
from django.db.models import Q, Prefetch
from django.utils import timezone
import json

from django.http import HttpResponse
from formulas.export import ExcelExporter
from datetime import datetime
from openpyxl.styles import Font

from formulas.models import (
    Formula, Category, Project, Variant, VariantValue,
    VariantGenerationRule, CalculationSession, SessionValue
)
from formulas.services import FormulaParser, FormulaCalculator, VariantGenerator, calculator


class CalculatorView(LoginRequiredMixin, View):
    """Главная страница калькулятора"""

    def get(self, request):
        # Только формулы текущего пользователя
        formulas = Formula.objects.filter(
            user=request.user
        ).select_related('category').order_by('category', 'order', 'symbol')

        categories = Category.objects.filter(user=request.user, parent=None)

        formulas_by_category = {}
        uncategorized = []

        for formula in formulas:
            if formula.category:
                cat_name = formula.category.name
                if cat_name not in formulas_by_category:
                    formulas_by_category[cat_name] = []
                formulas_by_category[cat_name].append(formula)
            else:
                uncategorized.append(formula)

        if uncategorized:
            formulas_by_category['Без категории'] = uncategorized

        context = {
            'formulas': formulas,
            'categories': categories,
            'formulas_by_category': formulas_by_category,
            'projects': Project.objects.filter(user=request.user)[:5],
        }
        return render(request, 'calculator/workspace.html', context)


class CalculateAllView(LoginRequiredMixin, View):
    """Страница вычисления ВСЕХ формул сразу"""

    def get(self, request):
        formulas = Formula.objects.filter(
            user=request.user
        ).select_related('category').order_by('category', 'order', 'symbol')

        input_formulas = [f for f in formulas if f.is_input]
        calculated_formulas = [f for f in formulas if not f.is_input]

        # Загружаем сохранённые значения из сессии
        session_values = request.session.get('calculator_values', {})

        context = {
            'input_formulas': input_formulas,
            'calculated_formulas': calculated_formulas,
            'all_formulas': formulas,
            'session_values': session_values,
        }
        return render(request, 'calculator/calculate_all.html', context)

    def post(self, request):
        """AJAX вычисление всех формул"""
        try:
            data = json.loads(request.body)
            input_values = {}

            # Преобразуем входные значения
            for key, val in data.get('values', {}).items():
                try:
                    if val != '' and val is not None:
                        input_values[key] = float(val)
                except (ValueError, TypeError):
                    pass

            # Сохраняем в сессию
            request.session['calculator_values'] = {k: str(v) for k, v in input_values.items()}

            # Вычисляем все формулы текущего пользователя
            all_formulas = list(Formula.objects.filter(user=request.user))
            results = calculator.calculate_all(all_formulas, input_values)

            # Форматируем ответ
            response_data = {
                'success': True,
                'results': {}
            }

            for symbol, result in results.items():
                response_data['results'][symbol] = {
                    'value': result.value,
                    'success': result.success,
                    'error': result.error,
                    'expression': result.formula_expression
                }

            return JsonResponse(response_data)

        except Exception as e:
            return JsonResponse({
                'success': False,
                'error': str(e)
            })


class FormulaCalculatorView(LoginRequiredMixin, View):
    """Страница вычисления конкретной формулы"""

    def get(self, request, pk):
        formula = get_object_or_404(Formula, pk=pk, user=request.user)
        all_formulas = list(Formula.objects.filter(user=request.user))

        # Определяем зависимости
        dependencies = []
        if not formula.is_input and formula.expression:
            dep_symbols = FormulaParser.extract_variables(formula.expression)
            formula_map = {f.symbol: f for f in all_formulas}

            def collect_deps(symbol, depth=0):
                if symbol in formula_map and depth < 10:
                    f = formula_map[symbol]
                    if f.symbol != formula.symbol:
                        dependencies.append(f)
                        if not f.is_input and f.expression:
                            for dep in FormulaParser.extract_variables(f.expression):
                                collect_deps(dep, depth + 1)

            for dep in dep_symbols:
                collect_deps(dep)

        # Уникальные зависимости
        seen = set()
        unique_deps = []
        for d in dependencies:
            if d.symbol not in seen:
                seen.add(d.symbol)
                unique_deps.append(d)

        # Загружаем сохранённые значения
        session_values = request.session.get('calculator_values', {})

        context = {
            'formula': formula,
            'dependencies': unique_deps,
            'all_formulas': all_formulas,
            'session_values': session_values,
        }
        return render(request, 'calculator/formula_calculate.html', context)


# ПРОЕКТЫ И ВАРИАНТЫ 

class ProjectListView(LoginRequiredMixin, ListView):
    """Список проектов"""
    model = Project
    template_name = 'calculator/project_list.html'
    context_object_name = 'projects'

    def get_queryset(self):
        return Project.objects.filter(user=self.request.user)


class ProjectCreateView(LoginRequiredMixin, CreateView):
    """Создание проекта"""
    model = Project
    template_name = 'calculator/project_form.html'
    fields = ['name', 'description']
    success_url = reverse_lazy('calculator:project_list')

    def get_form(self, form_class=None):
        form = super().get_form(form_class)
        form.fields['name'].widget.attrs.update({'class': 'form-control'})
        form.fields['description'].widget.attrs.update({'class': 'form-control', 'rows': 3})
        return form

    def form_valid(self, form):
        form.instance.user = self.request.user
        messages.success(self.request, 'Проект создан')
        return super().form_valid(form)


class ProjectDetailView(LoginRequiredMixin, DetailView):
    """Детали проекта с вариантами"""
    model = Project
    template_name = 'calculator/project_detail.html'
    context_object_name = 'project'

    def get_queryset(self):
        return Project.objects.filter(user=self.request.user)

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        
        # ИСПРАВЛЕНИЕ N+1: Prefetch related
        variants = self.object.variants.prefetch_related('values').all()
        
        # ФИЛЬТРАЦИЯ И ПОИСК
        search = self.request.GET.get('search', '').strip()
        filter_type = self.request.GET.get('filter', '')
        
        if search:
            variants = variants.filter(student_name__icontains=search)
        
        if filter_type == 'assigned':
            variants = variants.exclude(student_name='')
        elif filter_type == 'unassigned':
            variants = variants.filter(student_name='')
        
        context['variants'] = variants
        context['search_query'] = search
        context['current_filter'] = filter_type
        
        context['rules'] = self.object.generation_rules.select_related('formula')
        context['input_formulas'] = Formula.objects.filter(user=self.request.user, is_input=True)
        return context


class ProjectDeleteView(LoginRequiredMixin, DeleteView):
    """Удаление проекта"""
    model = Project
    template_name = 'calculator/project_confirm_delete.html'
    success_url = reverse_lazy('calculator:project_list')

    def get_queryset(self):
        return Project.objects.filter(user=self.request.user)


class GenerationRuleCreateView(LoginRequiredMixin, View):
    """Добавление правила генерации"""

    def post(self, request, project_id):
        project = get_object_or_404(Project, pk=project_id, user=request.user)

        formula_id = request.POST.get('formula')
        min_val = request.POST.get('min_value')
        max_val = request.POST.get('max_value')
        step = request.POST.get('step', 1)

        try:
            # Проверяем, что формула принадлежит пользователю
            formula = Formula.objects.get(pk=formula_id, user=request.user)
            
            # ВАЛИДАЦИЯ ПРАВИЛ
            min_value = float(min_val)
            max_value = float(max_val)
            step_value = float(step) if step else 1
            
            if not formula.is_input:
                messages.error(request, 'Можно создавать правила только для входных параметров')
                return redirect('calculator:project_detail', pk=project_id)
            
            if min_value > max_value:
                messages.error(request, 'Минимум не может быть больше максимума')
                return redirect('calculator:project_detail', pk=project_id)
            
            if step_value <= 0:
                messages.error(request, 'Шаг должен быть больше нуля')
                return redirect('calculator:project_detail', pk=project_id)

            VariantGenerationRule.objects.update_or_create(
                project=project,
                formula=formula,
                defaults={
                    'min_value': min_value,
                    'max_value': max_value,
                    'step': step_value,
                }
            )
            messages.success(request, f'Правило для {formula.symbol} добавлено')
        except Formula.DoesNotExist:
            messages.error(request, 'Формула не найдена')
        except ValueError:
            messages.error(request, 'Некорректные числовые значения')
        except Exception as e:
            messages.error(request, f'Ошибка: {e}')

        return redirect('calculator:project_detail', pk=project_id)


class GenerationRuleDeleteView(LoginRequiredMixin, View):
    """Удаление правила генерации"""

    def post(self, request, pk):
        rule = get_object_or_404(VariantGenerationRule, pk=pk, project__user=request.user)
        project_id = rule.project.id
        rule.delete()
        messages.success(request, 'Правило удалено')
        return redirect('calculator:project_detail', pk=project_id)


class GenerateVariantsView(LoginRequiredMixin, View):
    """Генерация вариантов"""

    def post(self, request, project_id):
        project = get_object_or_404(Project, pk=project_id, user=request.user)

        try:
            count = int(request.POST.get('count', 1))
            count = min(count, 100)  # Ограничение

            generator = VariantGenerator(project)
            variants = generator.generate_variants(count)

            messages.success(request, f'Создано {len(variants)} вариантов')
        except Exception as e:
            messages.error(request, f'Ошибка: {e}')

        return redirect('calculator:project_detail', pk=project_id)


class VariantDetailView(LoginRequiredMixin, DetailView):
    """Детали варианта"""
    model = Variant
    template_name = 'calculator/variant_detail.html'
    context_object_name = 'variant'

    def get_queryset(self):
        # ИСПРАВЛЕНИЕ N+1
        return Variant.objects.filter(
            project__user=self.request.user
        ).prefetch_related('values', 'values__formula', 'values__formula__category')

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        variant = self.object
        
        # Проверяем, изменились ли формулы
        formulas_outdated = False
        if variant.formulas_snapshot:
            current_formulas = {f.symbol: f for f in Formula.objects.filter(user=variant.project.user)}
            
            for symbol, snapshot_data in variant.formulas_snapshot.items():
                current = current_formulas.get(symbol)
                
                if not current:  # Формула удалена
                    formulas_outdated = True
                    break
                
                # Проверяем изменение выражения или типа
                if (current.expression != snapshot_data.get('expression') or
                    current.is_input != snapshot_data.get('is_input')):
                    formulas_outdated = True
                    break
        
        context['formulas_outdated'] = formulas_outdated
        context['last_recalculated_at'] = variant.last_recalculated_at
        
        # Группируем значения
        values = variant.values.all()
        input_values = [v for v in values if v.is_input]
        calculated_values = [v for v in values if not v.is_input]

        context['input_values'] = input_values
        context['calculated_values'] = calculated_values

        return context


class VariantUpdateStudentView(LoginRequiredMixin, View):
    """Обновление ФИО студента"""
    
    def post(self, request, pk):
        variant = get_object_or_404(Variant, pk=pk, project__user=request.user)
        
        student_name = request.POST.get('student_name', '').strip()
        variant.student_name = student_name
        variant.save()
        
        messages.success(request, 'ФИО студента обновлено')
        return redirect('calculator:variant_detail', pk=pk)


class VariantDeleteView(LoginRequiredMixin, View):
    """Удаление варианта"""

    def post(self, request, pk):
        variant = get_object_or_404(Variant, pk=pk, project__user=request.user)
        project_id = variant.project.id
        variant.delete()
        messages.success(request, 'Вариант удалён')
        return redirect('calculator:project_detail', pk=project_id)

class VariantRecalculateView(LoginRequiredMixin, View):
    """Пересчёт варианта"""

    def post(self, request, pk):
        variant = get_object_or_404(Variant, pk=pk, project__user=request.user)

        try:
            # Собираем входные значения
            input_values = {}
            for vv in variant.values.filter(is_input=True).select_related('formula'):
                input_values[vv.formula.symbol] = vv.value

            all_formulas = list(Formula.objects.filter(user=request.user).select_related('category'))

            # Вычисляем
            results = calculator.calculate_all(all_formulas, input_values)

            # Обновляем значения
            for symbol, result in results.items():
                current_formula = next((f for f in all_formulas if f.symbol == symbol), None)

                if current_formula and result.value is not None:
                    VariantValue.objects.update_or_create(
                        variant=variant,
                        formula=current_formula,
                        defaults={
                            'value': result.value,
                            'is_input': current_formula.is_input
                        }
                    )
            
            # Обновляем время пересчёта
            variant.last_recalculated_at = timezone.now()
            variant.save()

            messages.success(request, 'Вариант пересчитан по актуальным формулам')
        except Exception as e:
            messages.error(request, f'Ошибка при пересчёте: {e}')

        return redirect('calculator:variant_detail', pk=pk)

class ExportVariantView(LoginRequiredMixin, View):
    """Экспорт варианта в Excel"""

    def get(self, request, pk):
        variant = get_object_or_404(Variant, pk=pk, project__user=request.user)

        exporter = ExcelExporter()

        # Создаём лист с вариантом
        ws = exporter.wb.create_sheet(f"Вариант {variant.number}")

        # Информация
        ws.cell(row=1, column=1, value="Проект:").font = Font(bold=True)
        ws.cell(row=1, column=2, value=variant.project.name)
        ws.cell(row=2, column=1, value="Вариант:").font = Font(bold=True)
        ws.cell(row=2, column=2, value=f"№ {variant.number}")
        ws.cell(row=3, column=1, value="Студент:").font = Font(bold=True)
        ws.cell(row=3, column=2, value=variant.student_name or "—")
        ws.cell(row=4, column=1, value="Дата:").font = Font(bold=True)
        ws.cell(row=4, column=2, value=variant.created_at.strftime('%d.%m.%Y %H:%M'))

        # Заголовки таблицы
        headers = ["Символ", "Название", "Тип", "Значение", "Ед.изм."]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=6, column=col, value=header)
            exporter._apply_header_style(cell)

        # Данные
        values = variant.values.select_related('formula').order_by('-is_input', 'formula__symbol')
        for row_idx, vv in enumerate(values, 7):
            data = [
                vv.formula.symbol,
                vv.formula.name,
                "Входной" if vv.is_input else "Вычислено",
                round(vv.value, 4) if vv.value else "—",
                vv.formula.unit or "—"
            ]
            for col, val in enumerate(data, 1):
                cell = ws.cell(row=row_idx, column=col, value=val)
                exporter._apply_cell_style(cell, is_number=(col == 4), is_input=vv.is_input)

        exporter._auto_width(ws)

        filename = f"variant_{variant.number}_{datetime.now().strftime('%Y%m%d')}.xlsx"

        response = HttpResponse(
            exporter.get_file(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = f'attachment; filename="{filename}"'

        return response


class ExportProjectView(LoginRequiredMixin, View):
    """Экспорт проекта со всеми вариантами в один файл"""

    def get(self, request, pk):
        from formulas.export import export_full_data

        project = get_object_or_404(Project, pk=pk, user=request.user)
        variants = list(project.variants.prefetch_related('values', 'values__formula').order_by('number'))
        formulas = list(Formula.objects.filter(
            user=request.user
        ).select_related('category').order_by('-is_input', 'symbol'))

        file_data = export_full_data(
            formulas=formulas,
            project=project,
            variants=variants
        )

        # Безопасное имя файла
        safe_name = "".join(c for c in project.name if c.isalnum() or c in " _-")[:30]
        filename = f"project_{safe_name}_{datetime.now().strftime('%Y%m%d')}.xlsx"

        response = HttpResponse(
            file_data,
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = f'attachment; filename="{filename}"'

        return response